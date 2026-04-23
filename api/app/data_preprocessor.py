"""
資料預處理：將 Excel 指標 + 歷年 docx 意見表解析成 JSON 索引，供 LLM 使用。

執行方式（須先設定環境變數指向原始資料）：
    export INDICATORS_EXCEL=/path/to/指標.xlsx
    export FEWSHOT_DOCX_DIR=/path/to/委員評鑑結果資料目錄
    python3 -m app.data_preprocessor

輸出：
    api/data/indicators.json   — 指標索引（key: 年度_機構種類_代碼）
    api/data/fewshot.json      — few-shot 索引（key: 代碼）

注意：原始 Excel/docx 含評鑑委員個資與機構名稱，不納入版控。
      正常部署無需重新執行此腳本，api/data/*.json 已 bake 進 repo。
"""

import json
import os
import re
import sys
from collections import defaultdict
from pathlib import Path

DATA_DIR = Path(__file__).parent.parent / "data"


def _resolve_source_paths() -> tuple[Path, Path]:
    """從環境變數讀取原始資料路徑，缺一即報錯退出。"""
    excel = os.environ.get("INDICATORS_EXCEL")
    docx_dir = os.environ.get("FEWSHOT_DOCX_DIR")
    if not excel or not docx_dir:
        sys.exit(
            "錯誤：請設定環境變數 INDICATORS_EXCEL 與 FEWSHOT_DOCX_DIR 指向原始資料路徑。\n"
            "範例：\n"
            "  export INDICATORS_EXCEL=/path/to/指標.xlsx\n"
            "  export FEWSHOT_DOCX_DIR=/path/to/委員評鑑結果資料"
        )
    return Path(excel), Path(docx_dir)

YEAR_FOLDERS = {
    112: "112年委員評鑑結果資料",
    113: "113年委員評鑑結果資料",
    114: "114年委員評鑑結果資料",
}


def build_indicators(excel_path: Path) -> dict:
    """解析 Excel，建立指標索引。

    回傳結構：
    {
      "114_機構住宿式_A1": {
        "代碼": "A1",
        "指標種類": "經營管理效能",
        "指標內容": "...",
        "基準說明": "...",
        "評分標準": "..."
      },
      ...
    }
    """
    import openpyxl

    wb = openpyxl.load_workbook(excel_path, read_only=True)
    ws = wb["工作表1"]
    rows = list(ws.iter_rows(values_only=True))

    index = {}
    for row in rows[1:]:
        if not any(row):
            continue
        year, kind1, kind2, code, cat, cls, content, spec, method, scoring, *_ = row

        if not all([year, kind1, code, content]):
            continue

        # 機構種類 key（綜合式帶子類別）
        if kind1 == "綜合式" and kind2:
            type_key = f"{kind1}_{kind2}"
        else:
            type_key = str(kind1)

        key = f"{year}_{type_key}_{code}"
        index[key] = {
            "年度": year,
            "機構種類": type_key,
            "代碼": str(code),
            "指標種類": cat or "",
            "指標內容": str(content).strip(),
            "基準說明": str(spec).strip() if spec else "",
            "評分標準": str(scoring).strip() if scoring else "",
        }

    return index


def _extract_code(cell_text: str) -> str | None:
    """從儲存格文字提取指標代碼（A1, B12, C3, 1, 2...）。"""
    text = cell_text.strip()
    if re.fullmatch(r"[A-C]\d{1,2}", text):
        return text
    if re.fullmatch(r"\d{1,2}", text):
        return text
    return None


def build_fewshot(docx_base: Path) -> dict:
    """解析歷年 docx，建立 few-shot 意見索引。

    回傳結構：
    {
      "A1": [
        {"類型": "改善", "意見": "..."},
        {"類型": "建議", "意見": "..."},
        ...
      ],
      ...
    }
    最多保留每個代碼 30 筆，避免 JSON 過大。
    """
    from docx import Document

    index: dict[str, list] = defaultdict(list)
    MAX_PER_CODE = 30

    for year, folder_name in YEAR_FOLDERS.items():
        folder = docx_base / folder_name
        if not folder.exists():
            continue

        for fname in sorted(folder.iterdir()):
            if fname.suffix.lower() != ".docx":
                continue
            try:
                doc = Document(fname)
            except Exception:
                continue

            for table in doc.tables:
                current_type = None
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if not cells:
                        continue

                    # 判斷表格類型（改善 / 建議）
                    if "改善事項" in cells[0]:
                        current_type = "改善"
                        continue
                    if "建議事項" in cells[0]:
                        current_type = "建議"
                        continue

                    # 嘗試提取指標代碼
                    code = _extract_code(cells[0])
                    if not code or current_type is None:
                        continue

                    opinion = cells[2] if len(cells) > 2 else ""
                    opinion = opinion.strip()
                    if len(opinion) < 5:
                        continue

                    if len(index[code]) < MAX_PER_CODE:
                        index[code].append({
                            "類型": current_type,
                            "意見": opinion,
                            "年度": year,
                        })

    return dict(index)


def main():
    excel_path, docx_base = _resolve_source_paths()
    if not excel_path.exists():
        sys.exit(f"錯誤：INDICATORS_EXCEL 檔案不存在：{excel_path}")
    if not docx_base.exists():
        sys.exit(f"錯誤：FEWSHOT_DOCX_DIR 目錄不存在：{docx_base}")

    DATA_DIR.mkdir(parents=True, exist_ok=True)

    print("解析 Excel 指標...")
    indicators = build_indicators(excel_path)
    out_indicators = DATA_DIR / "indicators.json"
    out_indicators.write_text(json.dumps(indicators, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"  → {len(indicators)} 條指標 → {out_indicators}")

    print("解析歷年 docx few-shot...")
    fewshot = build_fewshot(docx_base)
    out_fewshot = DATA_DIR / "fewshot.json"
    out_fewshot.write_text(json.dumps(fewshot, ensure_ascii=False, indent=2), encoding="utf-8")
    total = sum(len(v) for v in fewshot.values())
    print(f"  → {len(fewshot)} 個代碼，共 {total} 筆意見 → {out_fewshot}")

    print("完成。")


if __name__ == "__main__":
    main()
